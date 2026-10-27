"""
FICHA FINANCEIRA e CONTROLE DA CONTRIBUIÇÃO DO SERVIDOR.

- Importação (`import_lines`): ficha financeira do sistema anterior (CSV/XLSX, leiaute `COLUMNS`), uma linha por
  rubrica × competência × matrícula. Valida tudo antes de gravar (`apply=False` desfaz). Repetir o arquivo atualiza
  o valor da mesma rubrica/competência (não duplica). A CLASSIFICAÇÃO diz qual rubrica é a contribuição ao RPPS e qual
  é a base; se vier em branco, o sistema SUGERE pelo nome (a conferir em `reclassify`).
- Ficha gerada (`build`): por ano, rubrica × mês, juntando os contracheques da Folha de Ativos do SIST e as linhas
  importadas (na competência que tem contracheque do SIST, vale o contracheque).
- Contribuições (`contributions`): por competência e vínculo — base, quanto foi DESCONTADO do servidor, quanto
  deveria ter sido (alíquota do parâmetro `folha.aliquota_segurado`), diferença, patronal estimada e a situação do
  REPASSE do órgão ao RPPS naquela competência (financeiro.Contribution do tipo SERVIDOR e receitas baixadas).
- `to_history`: leva as bases mensais ao histórico contributivo (salários de contribuição) sem alterar o que existe.
"""
import hashlib
import io
import re
import unicodedata
from collections import defaultdict
from datetime import date
from decimal import Decimal

from django.db import transaction
from django.db.models import Sum

from apps.audit.services import log_event
from apps.core.exceptions import BusinessRuleError
from apps.core.validators import only_digits

from .models import FinancialCategory, FinancialImport, FinancialLine

COLUMNS = [
    ("CPF", "Obrigatório."),
    ("MATRICULA", "Matrícula do servidor no órgão (liga a linha ao vínculo)."),
    ("COMPETENCIA", "Obrigatório. MM/AAAA."),
    ("TIPO_FOLHA", "MENSAL (padrão), DECIMO, FERIAS, RESCISAO ou COMPLEMENTAR."),
    ("CODIGO", "Obrigatório. Código da rubrica no sistema de origem."),
    ("RUBRICA", "Obrigatório. Nome da rubrica."),
    ("NATUREZA", "Obrigatório. P = provento, D = desconto, B = base/informativa."),
    ("CLASSIFICACAO", "CONTRIB_RPPS (desconto previdenciário do servidor), BASE_RPPS (base de contribuição), "
                      "PATRONAL (informativa), IRRF ou em branco. Em branco o sistema sugere pelo nome."),
    ("VALOR", "Obrigatório. Ex.: 1234,56."),
]
ZERO = Decimal("0.00")
KINDS = {"MENSAL", "DECIMO", "FERIAS", "RESCISAO", "COMPLEMENTAR"}


def _norm(txt: str) -> str:
    return unicodedata.normalize("NFKD", txt or "").encode("ascii", "ignore").decode().upper()


def suggest_category(name: str, nature: str) -> str:
    """Sugestão pelo nome da rubrica (sempre conferir): desconto previdenciário, base, patronal, IRRF."""
    n = _norm(name)
    if re.search(r"\bIRRF\b|IMPOSTO DE RENDA|\bI\.?R\.?\b", n):
        return "IRRF" if nature == "DESCONTO" else ""
    prev = re.search(r"PREV|RPPS|INSTITUTO|IPASE|CONTRIB|FUNDO", n) and not re.search(r"INSS|RGPS|COMPLEMENTAR|SAUDE|SINDIC", n)
    if not prev:
        return ""
    if "PATRONAL" in n:
        return "PATRONAL"
    if nature == "BASE" or re.search(r"\bBASE\b", n):
        return "BASE_RPPS"
    return "CONTRIB_RPPS" if nature == "DESCONTO" else ""


def _comp(v) -> date:
    from datetime import datetime

    if isinstance(v, datetime):
        return v.date().replace(day=1)
    if isinstance(v, date):
        return v.replace(day=1)
    s = str(v or "").strip()
    m = re.fullmatch(r"(\d{1,2})[/\-.](\d{4})", s) or re.fullmatch(r"(\d{4})[/\-.](\d{1,2})(?:[/\-.]\d{1,2})?", s)
    if not m:
        m2 = re.fullmatch(r"\d{1,2}/(\d{1,2})/(\d{4})", s)
        if m2:
            return date(int(m2.group(2)), int(m2.group(1)), 1)
        raise ValueError(f"competência inválida: {s!r} (use MM/AAAA)")
    a, b = m.groups()
    month, year = (int(a), int(b)) if len(a) <= 2 else (int(b), int(a))
    if not 1 <= month <= 12 or year < 1990:
        raise ValueError(f"competência inválida: {s!r}")
    return date(year, month, 1)


@transaction.atomic
def import_lines(institution, content: bytes, file_name: str, *, person=None, apply=False) -> dict:
    from apps.employment.models import EmploymentLink
    from apps.people.importers.sprev import _dec, _s, read_table
    from apps.people.models import Person

    rows = read_table(content, file_name)
    if not rows:
        raise BusinessRuleError("Arquivo vazio.")
    missing = {"CPF", "COMPETENCIA", "CODIGO", "RUBRICA", "NATUREZA", "VALOR"} - set(rows[0].keys())
    if missing:
        raise BusinessRuleError(f"Colunas obrigatórias ausentes: {', '.join(sorted(missing))}. Baixe a tabela padrão.")
    batch = FinancialImport.objects.create(institution=institution, file_name=file_name[:255], person=person,
                                           sha256=hashlib.sha256(content).hexdigest())
    people, links = {}, {}
    created = updated = 0
    errors, rubrics, comps, servers = [], {}, set(), set()
    for i, row in enumerate(rows, start=2):
        try:
            cpf = only_digits(_s(row.get("CPF"))).zfill(11)
            if cpf not in people:
                people[cpf] = Person.objects.filter(institution=institution, cpf=cpf).first()
            p = people[cpf]
            if p is None:
                raise ValueError("CPF não está no cadastro de servidores")
            if person is not None and p.pk != person.pk:
                raise ValueError("linha de outro servidor (esta importação é só da pasta aberta)")
            reg = _s(row.get("MATRICULA"))
            reg = reg[:-2] if reg.endswith(".0") else reg
            key = (p.pk, reg)
            if key not in links:
                qs = EmploymentLink.objects.filter(person=p)
                links[key] = (qs.filter(registration=reg).first() if reg else None) or (qs.first() if qs.count() == 1 else None)
            comp = _comp(row.get("COMPETENCIA"))
            kind = (_s(row.get("TIPO_FOLHA")).upper() or "MENSAL")
            kind = {"13": "DECIMO", "13O": "DECIMO", "DECIMO TERCEIRO": "DECIMO"}.get(kind, kind)
            if kind not in KINDS:
                raise ValueError(f"TIPO_FOLHA inválido: {kind}")
            code, name = _s(row.get("CODIGO"))[:20], _s(row.get("RUBRICA"))[:120]
            if not code or not name:
                raise ValueError("CODIGO e RUBRICA são obrigatórios")
            nat = {"P": "PROVENTO", "D": "DESCONTO", "B": "BASE"}.get(_s(row.get("NATUREZA")).upper()[:1])
            if nat is None:
                raise ValueError("NATUREZA deve ser P, D ou B")
            amount = _dec(row.get("VALOR"))
            if amount is None:
                raise ValueError("VALOR não informado")
            cat = _s(row.get("CLASSIFICACAO")).upper()
            if cat and cat not in FinancialCategory.values:
                raise ValueError(f"CLASSIFICACAO inválida: {cat}")
            cat = cat or suggest_category(name, nat)
            _, c = FinancialLine.objects.update_or_create(
                person=p, registration=reg, competence=comp, payroll_kind=kind, code=code,
                defaults={"institution": institution, "link": links[key], "name": name, "nature": nat,
                          "category": cat, "amount": abs(amount), "batch": batch})
            created += int(c)
            updated += int(not c)
            rubrics.setdefault(code, {"codigo": code, "rubrica": name, "natureza": nat, "classificacao": cat})
            comps.add(comp)
            servers.add(p.pk)
        except (ValueError, BusinessRuleError) as exc:
            if len(errors) < 300:
                errors.append({"linha": i, "mensagem": str(exc)})
    summary = {"linhas": len(rows), "incluidas": created, "atualizadas": updated, "erros": errors,
               "total_erros": len(rows) - created - updated, "servidores": len(servers),
               "periodo": [min(comps).strftime("%m/%Y"), max(comps).strftime("%m/%Y")] if comps else None,
               "rubricas": sorted(rubrics.values(), key=lambda r: r["codigo"]),
               "sem_classificacao_rpps": not any(r["classificacao"] == "CONTRIB_RPPS" for r in rubrics.values())}
    if not apply:
        transaction.set_rollback(True)
        return {**summary, "gravado": False}
    batch.summary = {k: v for k, v in summary.items() if k != "erros"}
    batch.save(update_fields=["summary"])
    log_event("CREATE", module="ativos.ficha_financeira", obj=batch,
              object_repr=f"Ficha financeira {file_name}: {created} incluída(s), {updated} atualizada(s)")
    return {**summary, "gravado": True}


@transaction.atomic
def reclassify(institution, code: str, category: str, person=None) -> int:
    if category not in FinancialCategory.values:
        raise BusinessRuleError("Classificação inválida.")
    qs = FinancialLine.objects.filter(institution=institution, code=code)
    if person is not None:
        qs = qs.filter(person=person)
    n = qs.update(category=category)
    log_event("UPDATE", module="ativos.ficha_financeira",
              object_repr=f"Rubrica {code} da ficha financeira classificada como '{category or 'Outras'}' ({n} linha(s))")
    return n


def rubrics(institution, person=None) -> list[dict]:
    qs = FinancialLine.objects.filter(institution=institution)
    if person is not None:
        qs = qs.filter(person=person)
    out = {}
    for r in qs.values("code", "name", "nature", "category").annotate(total=Sum("amount")).order_by("code"):
        out.setdefault(r["code"], {"codigo": r["code"], "rubrica": r["name"], "natureza": r["nature"],
                                   "classificacao": r["category"], "total": str(r["total"])})
    return list(out.values())


# ---------------------------------------------------------------- fontes unificadas
def _folha_rows(person, year=None):
    """Contracheques da Folha de Ativos (exceto folhas ABERTAS) como linhas no mesmo formato da ficha importada."""
    from .models import ActivePayslip, RubricCalc

    qs = ActivePayslip.objects.filter(person=person).exclude(run__status="ABERTA").select_related("run", "run__organ", "link")
    if year:
        qs = qs.filter(run__year=year)
    out = []
    for s in qs.prefetch_related("items__rubric"):
        comp = date(s.run.year, s.run.month, 1)
        base = {"comp": comp, "kind": "MENSAL", "link": s.link, "reg": s.registration, "fonte": "FOLHA",
                "situacao": s.run.get_status_display()}
        for it in s.items.all():
            cat = ("CONTRIB_RPPS" if it.rubric and it.rubric.calc == RubricCalc.CONTRIB_RPPS
                   else "IRRF" if it.rubric and it.rubric.calc == RubricCalc.IRRF else "")
            out.append({**base, "code": it.code, "name": it.name, "nature": it.kind, "category": cat, "amount": it.amount})
        out.append({**base, "code": "B-RPPS", "name": "BASE DE CONTRIBUIÇÃO RPPS", "nature": "BASE",
                    "category": "BASE_RPPS", "amount": s.rpps_base})
        out.append({**base, "code": "B-PATR", "name": "CONTRIBUIÇÃO PATRONAL (informativa)", "nature": "BASE",
                    "category": "PATRONAL", "amount": s.employer_contribution})
    return out


def _history_rows(person, year, covered):
    """Salários de contribuição RPPS (histórico contributivo) das competências sem ficha: base sem desconto informado."""
    from apps.previdencia.models import ContributionSalary

    qs = ContributionSalary.objects.filter(person=person, regime="RPPS")
    if year:
        qs = qs.filter(competence__year=year)
    out = []
    for cs in qs:
        if cs.competence in covered:
            continue
        out.append({"comp": cs.competence, "kind": "MENSAL", "link": None, "reg": "", "fonte": "HISTORICO",
                    "situacao": "Histórico contributivo", "code": "HIST", "name": "SALÁRIO DE CONTRIBUIÇÃO (histórico)",
                    "nature": "BASE", "category": "BASE_RPPS", "amount": cs.amount, "orgao": cs.source_label})
    return out


def _all_rows(person, year=None):
    rows = _folha_rows(person, year)
    have = {(r["reg"], r["comp"]) for r in rows}
    qs = FinancialLine.objects.filter(person=person).select_related("link", "link__organ")
    if year:
        qs = qs.filter(competence__year=year)
    for ln in qs:
        if ln.payroll_kind == "MENSAL" and (ln.registration, ln.competence) in have:
            continue  # competência já calculada pela Folha de Ativos do SIST
        rows.append({"comp": ln.competence, "kind": ln.payroll_kind, "link": ln.link, "reg": ln.registration,
                     "fonte": "IMPORTADA", "situacao": "Importada", "code": ln.code, "name": ln.name,
                     "nature": ln.nature, "category": ln.category, "amount": ln.amount})
    return rows


def years(person) -> list[int]:
    from .models import ActivePayslip

    ys = set(FinancialLine.objects.filter(person=person).values_list("competence__year", flat=True))
    ys |= set(ActivePayslip.objects.filter(person=person).values_list("run__year", flat=True))
    return sorted(ys, reverse=True)


def build(person, year: int, registration: str | None = None) -> dict:
    """Ficha financeira do ano: rubrica × mês (13 = 13º salário; férias/rescisão/complementar somam no mês)."""
    rows = [r for r in _all_rows(person, year) if registration in (None, "", r["reg"])]
    lines = {}
    tot = defaultdict(lambda: defaultdict(lambda: ZERO))
    for r in rows:
        col = 13 if r["kind"] == "DECIMO" else r["comp"].month
        key = (r["nature"], r["code"], r["name"])
        ln = lines.setdefault(key, {"codigo": r["code"], "rubrica": r["name"], "natureza": r["nature"],
                                    "classificacao": r["category"], "valores": defaultdict(lambda: ZERO)})
        ln["valores"][col] += r["amount"]
        if r["nature"] == "PROVENTO":
            tot["proventos"][col] += r["amount"]
        elif r["nature"] == "DESCONTO":
            tot["descontos"][col] += r["amount"]
        if r["category"] in ("CONTRIB_RPPS", "BASE_RPPS"):
            tot[r["category"]][col] += r["amount"]
    order = {"PROVENTO": 0, "DESCONTO": 1, "BASE": 2}
    out_lines = []
    for (nat, code, _), ln in sorted(lines.items(), key=lambda kv: (order[kv[0][0]], kv[0][1])):
        vals = {str(m): str(v) for m, v in ln["valores"].items()}
        out_lines.append({**ln, "valores": vals, "total": str(sum(ln["valores"].values(), ZERO))})
    months = range(1, 14)
    totals = {k: {str(m): str(tot[k][m]) for m in months if tot[k][m]} for k in ("proventos", "descontos", "CONTRIB_RPPS", "BASE_RPPS")}
    totals["liquido"] = {str(m): str(tot["proventos"][m] - tot["descontos"][m]) for m in months
                         if tot["proventos"][m] or tot["descontos"][m]}
    for k in list(totals):
        totals[k]["total"] = str(sum((Decimal(v) for v in totals[k].values()), ZERO))
    fontes = sorted({r["fonte"] for r in rows})
    regs = sorted({r["reg"] for r in _all_rows(person, year) if r["reg"]})
    return {"ano": year, "matricula": registration or "", "matriculas": regs, "fontes": fontes, "linhas": out_lines,
            "totais": totals, "anos": years(person)}


def export_xlsx(person, year: int, registration: str | None = None) -> bytes:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    data = build(person, year, registration)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Ficha {year}"
    ws.append([f"FICHA FINANCEIRA {year} — {person.name} — CPF {person.cpf[:3]}.{person.cpf[3:6]}.{person.cpf[6:9]}-{person.cpf[9:]}"
               + (f" — matrícula {registration}" if registration else "")])
    ws["A1"].font = Font(bold=True, size=12)
    head = ["Código", "Rubrica", "Natureza"] + [f"{m:02d}" for m in range(1, 13)] + ["13º", "Total"]
    ws.append([])
    ws.append(head)
    for c in ws[3]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="9F1D2B")
        c.alignment = Alignment(horizontal="center")
    nat = {"PROVENTO": "Provento", "DESCONTO": "Desconto", "BASE": "Base"}

    def num(v):
        return Decimal(v) if v not in (None, "") else None

    for ln in data["linhas"]:
        ws.append([ln["codigo"], ln["rubrica"], nat[ln["natureza"]]]
                  + [num(ln["valores"].get(str(m))) for m in range(1, 14)] + [num(ln["total"])])
    ws.append([])
    for key, label in (("proventos", "TOTAL DE PROVENTOS"), ("descontos", "TOTAL DE DESCONTOS"), ("liquido", "LÍQUIDO"),
                       ("BASE_RPPS", "BASE RPPS"), ("CONTRIB_RPPS", "CONTRIBUIÇÃO DO SERVIDOR AO RPPS")):
        t = data["totais"].get(key, {})
        ws.append(["", label, ""] + [num(t.get(str(m))) for m in range(1, 14)] + [num(t.get("total"))])
        for c in ws[ws.max_row]:
            c.font = Font(bold=True)
    for row in ws.iter_rows(min_row=4, min_col=4):
        for c in row:
            c.number_format = "#,##0.00"
    ws.column_dimensions["B"].width = 42
    for col in "DEFGHIJKLMNOPQ":
        ws.column_dimensions[col].width = 12
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ---------------------------------------------------------------- controle da contribuição
def contributions(person, year: int | None = None) -> dict:
    """Por competência e matrícula: base, descontado, devido pela alíquota, diferença, patronal e repasse do órgão."""
    from apps.financeiro.models import Contribution, Revenue
    from apps.institutions.services import get_parameter

    inst = person.institution
    acc = {}
    for r in _all_rows(person, year):
        if r["kind"] not in ("MENSAL", "DECIMO", "COMPLEMENTAR", "RESCISAO", "FERIAS"):
            continue
        k = (r["comp"], r["kind"], r["reg"])
        a = acc.setdefault(k, {"base": ZERO, "desc": ZERO, "patronal": None, "fonte": r["fonte"], "link": r["link"],
                               "situacao": r["situacao"]})
        if r["category"] == "BASE_RPPS":
            a["base"] += r["amount"]
        elif r["category"] == "CONTRIB_RPPS":
            a["desc"] += r["amount"]
        elif r["category"] == "PATRONAL":
            a["patronal"] = (a["patronal"] or ZERO) + r["amount"]
    covered = {k[0] for k in acc}
    for r in _history_rows(person, year, covered):
        acc[(r["comp"], "MENSAL", "")] = {"base": r["amount"], "desc": ZERO, "patronal": None, "fonte": "HISTORICO",
                                         "link": None, "situacao": r["situacao"], "orgao": r["orgao"]}
    rate_cache, items = {}, []
    tot = defaultdict(lambda: ZERO)
    for (comp, kind, reg), a in sorted(acc.items(), key=lambda kv: (kv[0][0], kv[0][1]), reverse=True):
        if not a["base"] and not a["desc"]:
            continue
        if comp not in rate_cache:
            seg = get_parameter(inst, "folha.aliquota_segurado", comp)
            pat = get_parameter(inst, "arrecadacao.aliquota_patronal", comp)
            rate_cache[comp] = (Decimal(str(seg)) if seg not in (None, "") else None,
                                Decimal(str(pat)) if pat not in (None, "") else None)
        rate, prate = rate_cache[comp]
        avisos = []
        base, base_estimada = a["base"], False
        if not base and a["desc"] and rate:
            base, base_estimada = (a["desc"] / rate).quantize(Decimal("0.01")), True
            avisos.append("base estimada pelo desconto")
        devido = (base * rate).quantize(Decimal("0.01")) if rate is not None else None
        dif = (a["desc"] - devido) if devido is not None else None
        if rate is None:
            avisos.append("alíquota do servidor não cadastrada para a competência")
        elif not a["desc"]:
            avisos.append("desconto não informado (só a base no histórico)" if a["fonte"] == "HISTORICO"
                          else "nenhum desconto previdenciário encontrado")
        elif abs(dif) > Decimal("0.05"):
            avisos.append("descontado a menor" if dif < 0 else "descontado a maior")
        patronal = a["patronal"] if a["patronal"] is not None else (
            (base * prate).quantize(Decimal("0.01")) if prate is not None else None)
        repasse = {"situacao": "Sem guia lançada", "devido": None, "pago": None}
        link = a["link"]
        if link is not None:
            guias = Contribution.objects.filter(institution=inst, organ_id=link.organ_id, competence=comp,
                                                kind="SERVIDOR", active=True)
            if guias.exists():
                due = guias.aggregate(s=Sum("amount"))["s"] or ZERO
                paid = Revenue.objects.filter(contribution__in=guias, active=True).aggregate(s=Sum("amount"))["s"] or ZERO
                repasse = {"situacao": "Quitado" if paid >= due else "Pago em parte" if paid else "Em aberto",
                           "devido": str(due), "pago": str(paid)}
        items.append({
            "competencia": comp.strftime("%m/%Y"), "folha": kind, "matricula": reg,
            "orgao": link.organ.name if link is not None else a.get("orgao", ""), "fonte": a["fonte"],
            "base": str(base), "base_estimada": base_estimada, "descontado": str(a["desc"]),
            "aliquota_aplicada": str((a["desc"] / base * 100).quantize(Decimal("0.01"))) if base else None,
            "aliquota_legal": str((rate * 100).quantize(Decimal("0.01"))) if rate is not None else None,
            "devido": str(devido) if devido is not None else None, "diferenca": str(dif) if dif is not None else None,
            "patronal": str(patronal) if patronal is not None else None, "repasse_orgao": repasse, "avisos": avisos,
        })
        tot["base"] += base
        tot["descontado"] += a["desc"]
        tot["devido"] += devido or ZERO
        tot["patronal"] += patronal or ZERO
    tot["diferenca"] = tot["descontado"] - tot["devido"]
    return {"competencias": items, "totais": {k: str(v) for k, v in tot.items()},
            "com_aviso": sum(1 for i in items if i["avisos"]), "anos": years(person)}


@transaction.atomic
def to_history(person) -> dict:
    """Bases mensais (folha MENSAL) → salários de contribuição RPPS. Não altera competências já registradas."""
    from apps.previdencia.models import ContributionSalary

    by_comp = defaultdict(lambda: ZERO)
    for r in _all_rows(person):
        if r["kind"] == "MENSAL" and r["category"] == "BASE_RPPS":
            by_comp[(r["comp"], r["link"].organ.name if r["link"] else "")] += r["amount"]
    data = contributions(person)
    for c in data["competencias"]:  # base estimada pelo desconto quando a ficha não traz a base
        if c["folha"] == "MENSAL" and c["base_estimada"]:
            m, y = c["competencia"].split("/")
            by_comp.setdefault((date(int(y), int(m), 1), c["orgao"]), Decimal(c["base"]))
    created = existing = 0
    for (comp, organ), amount in sorted(by_comp.items(), key=lambda kv: kv[0][0]):
        if comp < date(1994, 7, 1) or amount <= 0:
            continue
        if ContributionSalary.objects.filter(person=person, competence=comp, regime="RPPS").exists():
            existing += 1
            continue
        ContributionSalary.objects.create(institution=person.institution, person=person, competence=comp, regime="RPPS",
                                          source_label=organ[:150], amount=amount, origin="FICHA_FINANCEIRA")
        created += 1
    log_event("UPDATE", module="ativos.ficha_financeira", obj=person,
              object_repr=f"{person.name}: {created} salário(s) de contribuição incluído(s) a partir da ficha financeira")
    return {"incluidos": created, "ja_existiam": existing}


def template_xlsx() -> bytes:
    import openpyxl
    from openpyxl.styles import Font, PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "FICHA_FINANCEIRA"
    ws.append([c for c, _ in COLUMNS])
    required = {"CPF", "COMPETENCIA", "CODIGO", "RUBRICA", "NATUREZA", "VALOR"}
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="C2410C" if cell.value in required else "1E3A5F")
        ws.column_dimensions[cell.column_letter].width = max(14, len(str(cell.value)) + 4)
    for col in range(1, len(COLUMNS) + 1):
        for row in range(2, 1001):
            ws.cell(row=row, column=col).number_format = "@"
    ws.freeze_panes = "A2"
    ins = wb.create_sheet("INSTRUCOES")
    ins.append(["Coluna", "Como preencher"])
    for c, h in COLUMNS:
        ins.append([c, h])
    ins.append([])
    ins.append(["Exemplo", "52998224725;1234;03/2019;MENSAL;001;VENCIMENTO;P;;2500,00"])
    ins.append(["Exemplo", "52998224725;1234;03/2019;MENSAL;501;CONTRIBUIÇÃO PREVIDENCIÁRIA;D;CONTRIB_RPPS;275,00"])
    ins.append(["Exemplo", "52998224725;1234;03/2019;MENSAL;900;BASE PREVIDÊNCIA;B;BASE_RPPS;2500,00"])
    ins.append(["Dica", "Uma linha por rubrica de cada mês. Repetir o arquivo atualiza os valores (não duplica)."])
    ins.column_dimensions["A"].width = 18
    ins.column_dimensions["B"].width = 120
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@transaction.atomic
def add_manual(person, *, competence, link=None, base, rate=None, withheld=None, kind="MENSAL", to_history_too=True) -> dict:
    """
    Lança uma competência no controle da contribuição sem arquivo: base e valor descontado (se o desconto não for
    informado, usa base × alíquota informada ou a vigente). Fica como ficha financeira de origem "Lançamento manual".
    """
    from apps.institutions.services import get_parameter
    from apps.previdencia.models import ContributionSalary

    competence = competence.replace(day=1)
    if base is None or base <= 0:
        raise BusinessRuleError("Informe a base de contribuição.")
    if kind not in KINDS:
        raise BusinessRuleError("Tipo de folha inválido.")
    if rate is None:
        v = get_parameter(person.institution, "folha.aliquota_segurado", competence)
        if v in (None, "") and withheld is None:
            raise BusinessRuleError(f"Não há alíquota do servidor cadastrada para {competence:%m/%Y}: informe a alíquota "
                                    "ou o valor descontado (ou cadastre em Arrecadação › Alíquotas).")
        rate = Decimal(str(v)) if v not in (None, "") else None
    if withheld is None:
        withheld = (base * rate).quantize(Decimal("0.01"))
    batch = FinancialImport.objects.create(institution=person.institution, file_name="Lançamento manual",
                                           sha256="", person=person)
    reg = link.registration if link is not None else ""
    for code, name, nature, cat, amount in (("BASE-RPPS", "BASE DE CONTRIBUIÇÃO RPPS", "BASE", "BASE_RPPS", base),
                                            ("DESC-RPPS", "CONTRIBUIÇÃO DO SERVIDOR AO RPPS", "DESCONTO", "CONTRIB_RPPS", withheld)):
        FinancialLine.objects.update_or_create(
            person=person, registration=reg, competence=competence, payroll_kind=kind, code=code,
            defaults={"institution": person.institution, "link": link, "name": name, "nature": nature,
                      "category": cat, "amount": amount, "batch": batch})
    hist = False
    if to_history_too and kind == "MENSAL" and competence >= date(1994, 7, 1) and not ContributionSalary.objects.filter(
            person=person, competence=competence, regime="RPPS").exists():
        ContributionSalary.objects.create(institution=person.institution, person=person, competence=competence,
                                          regime="RPPS", amount=base, origin="FICHA_FINANCEIRA",
                                          source_label=(link.organ.name if link is not None else "")[:150])
        hist = True
    log_event("CREATE", module="ativos.ficha_financeira", obj=person,
              object_repr=f"{person.name}: contribuição {competence:%m/%Y} lançada (base {base}, descontado {withheld})")
    return {"competencia": competence.strftime("%m/%Y"), "base": str(base), "descontado": str(withheld),
            "historico": hist}
