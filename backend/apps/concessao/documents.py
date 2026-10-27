"""
Minutas e portarias da concessão em Word (.docx).

O texto vem de modelos editáveis (DocumentTemplate) com marcadores entre chaves, ex.: {nome}, {valor}. Se a
instituição não tiver modelo próprio, usa-se o modelo padrão abaixo — redação sugerida, a ser ajustada pela
Procuradoria Jurídica da instituição em Concessão › Modelos de documentos.

Cada documento gerado é guardado com hash SHA-256 e anexado ao processo do Protocolo, onde o setor de
Benefícios o encontra automaticamente (além do aviso na Central de Alertas).
"""
import hashlib
import io
import re
from decimal import ROUND_HALF_UP, Decimal

from django.core.files.base import ContentFile
from django.utils import timezone

from .models import ConcessionDocument, ConcessionKind, DocumentKind, DocumentTemplate

MONTHS = ["janeiro", "fevereiro", "março", "abril", "maio", "junho", "julho", "agosto", "setembro", "outubro",
          "novembro", "dezembro"]

# ------------------------------------------------------------------ valor por extenso
_UNITS = ["zero", "um", "dois", "três", "quatro", "cinco", "seis", "sete", "oito", "nove", "dez", "onze", "doze",
          "treze", "quatorze", "quinze", "dezesseis", "dezessete", "dezoito", "dezenove"]
_TENS = ["", "", "vinte", "trinta", "quarenta", "cinquenta", "sessenta", "setenta", "oitenta", "noventa"]
_HUNDREDS = ["", "cento", "duzentos", "trezentos", "quatrocentos", "quinhentos", "seiscentos", "setecentos",
             "oitocentos", "novecentos"]


def _below_1000(n: int) -> str:
    if n == 100:
        return "cem"
    parts = []
    h, rest = divmod(n, 100)
    if h:
        parts.append(_HUNDREDS[h])
    if rest:
        if rest < 20:
            parts.append(_UNITS[rest])
        else:
            t, u = divmod(rest, 10)
            parts.append(_TENS[t] + (f" e {_UNITS[u]}" if u else ""))
    return " e ".join(parts)


def _integer_words(n: int) -> str:
    if n == 0:
        return "zero"
    groups = []
    scales = [("", ""), ("mil", "mil"), ("milhão", "milhões"), ("bilhão", "bilhões")]
    i = 0
    while n > 0:
        n, g = divmod(n, 1000)
        if g:
            if i == 1:
                words = "mil" if g == 1 else f"{_below_1000(g)} mil"
            elif i > 1:
                words = f"{_below_1000(g)} {scales[i][0] if g == 1 else scales[i][1]}"
            else:
                words = _below_1000(g)
            groups.append((g, words))
        i += 1
    groups.reverse()
    out = ""
    for idx, (g, words) in enumerate(groups):
        if idx == 0:
            out = words
        else:
            last = idx == len(groups) - 1
            sep = " e " if last and (g < 100 or g % 100 == 0) else ", "
            out += sep + words
    return out


def reais_por_extenso(value) -> str:
    v = Decimal(value).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)
    inteiro, cent = int(v), int((v - int(v)) * 100)
    parts = []
    if inteiro:
        w = _integer_words(inteiro)
        de = " de" if inteiro >= 1_000_000 and inteiro % 1_000_000 == 0 else ""
        parts.append(f"{w}{de} {'real' if inteiro == 1 else 'reais'}")
    if cent:
        parts.append(f"{_integer_words(cent)} {'centavo' if cent == 1 else 'centavos'}")
    return " e ".join(parts) or "zero real"


def br_money(v) -> str:
    s = f"{Decimal(v):,.2f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def br_date(d) -> str:
    return f"{d.day:02d}/{d.month:02d}/{d.year}" if d else "____/____/______"


def long_date(d) -> str:
    return f"{d.day} de {MONTHS[d.month - 1]} de {d.year}" if d else "____ de __________ de ______"


# ------------------------------------------------------------------ modelos padrão
DEFAULT_TEMPLATES = {
    (DocumentKind.PORTARIA, ConcessionKind.APOSENTADORIA): (
        "PORTARIA {sigla} Nº {numero_portaria}, DE {data_portaria_extenso}",
        "O PRESIDENTE DO {instituicao_maiusculo} – {sigla}, no uso de suas atribuições legais, e tendo em vista o que "
        "consta do Processo Administrativo nº {processo},\n"
        "RESOLVE:\n"
        "Art. 1º CONCEDER APOSENTADORIA {modalidade_maiusculo} ao(à) servidor(a) {nome}, CPF {cpf}, matrícula nº "
        "{matricula}, ocupante do cargo de {cargo}, lotado(a) em {orgao}, com fundamento no {fundamento} e na "
        "{fundamento_local}, a partir de {data_validade}.\n"
        "Art. 2º Os proventos mensais ficam fixados em R$ {valor} ({valor_extenso}), calculados com {forma_proventos}, "
        "{paridade}, conforme demonstrativo de fixação de proventos constante dos autos.\n"
        "Art. 3º Esta Portaria entra em vigor na data de sua publicação, produzindo efeitos a partir de {data_validade}."),
    (DocumentKind.PORTARIA, ConcessionKind.PENSAO): (
        "PORTARIA {sigla} Nº {numero_portaria}, DE {data_portaria_extenso}",
        "O PRESIDENTE DO {instituicao_maiusculo} – {sigla}, no uso de suas atribuições legais, e tendo em vista o que "
        "consta do Processo Administrativo nº {processo},\n"
        "RESOLVE:\n"
        "Art. 1º CONCEDER PENSÃO POR MORTE aos dependentes do(a) ex-servidor(a) {instituidor}, CPF {cpf_instituidor}, "
        "falecido(a) em {data_obito}, com fundamento no {fundamento}, a partir de {data_validade}, na forma abaixo:\n"
        "{cotas}\n"
        "Art. 2º O valor total da pensão fica fixado em R$ {valor} ({valor_extenso}), conforme demonstrativo de cálculo "
        "constante dos autos.\n"
        "Art. 3º Esta Portaria entra em vigor na data de sua publicação, produzindo efeitos a partir de {data_validade}."),
    (DocumentKind.PORTARIA, ConcessionKind.REVISAO): (
        "PORTARIA {sigla} Nº {numero_portaria}, DE {data_portaria_extenso}",
        "O PRESIDENTE DO {instituicao_maiusculo} – {sigla}, no uso de suas atribuições legais, e tendo em vista o que "
        "consta do Processo Administrativo nº {processo},\n"
        "RESOLVE:\n"
        "Art. 1º REVISAR o benefício de {beneficio_revisto} de {nome}, CPF {cpf}, fixando o valor mensal em R$ {valor} "
        "({valor_extenso}), em substituição ao valor de R$ {valor_anterior}, a partir de {data_validade}.\n"
        "Art. 2º Motivo da revisão: {motivo_revisao}\n"
        "Art. 3º Esta Portaria entra em vigor na data de sua publicação, produzindo efeitos a partir de {data_validade}."),
    (DocumentKind.CIENCIA_INDEFERIMENTO, ""): (
        "OFÍCIO {sigla} – CIÊNCIA DE INDEFERIMENTO",
        "Ao(À) Sr(a). {nome}, CPF {cpf}.\n"
        "Comunicamos que o requerimento de {tipo_concessao} objeto do Processo Administrativo nº {processo} foi "
        "INDEFERIDO pela Presidência do {sigla}, pelos seguintes fundamentos:\n"
        "{motivo_indeferimento}\n"
        "Fica facultada a apresentação de recurso no prazo legal.\n"
        "Barra Mansa, {hoje_extenso}.\n"
        "Termo de ciência: ______________________________________   Data: ____/____/______"),
    (DocumentKind.PORTARIA_ANULACAO, ""): (
        "PORTARIA {sigla} Nº {numero_portaria}, DE {data_portaria_extenso}",
        "O PRESIDENTE DO {instituicao_maiusculo} – {sigla}, no uso de suas atribuições legais, em cumprimento à "
        "decisão do Tribunal de Contas do Estado do Rio de Janeiro no processo nº {tce_processo}, que julgou ILEGAL o "
        "ato de concessão,\n"
        "RESOLVE:\n"
        "Art. 1º ANULAR a {portaria_original}, que concedeu {tipo_concessao} a {nome}, CPF {cpf}.\n"
        "Art. 2º Determinar a exclusão do benefício da folha de pagamento do {sigla}"
        "{reimplantacao}.\n"
        "Art. 3º Esta Portaria entra em vigor na data de sua publicação."),
    (DocumentKind.TERMO_DESISTENCIA, ""): (
        "TERMO DE DESISTÊNCIA E ARQUIVAMENTO",
        "Processo Administrativo nº {processo} – requerimento de {tipo_concessao} de {nome}, CPF {cpf}, "
        "matrícula {matricula}.\n"
        "O(A) requerente manifestou, por escrito, a DESISTÊNCIA do requerimento acima, pelos seguintes motivos:\n"
        "{motivo_desistencia}\n"
        "Em razão da desistência, o processo é arquivado sem a concessão do benefício. A desistência não impede novo "
        "requerimento a qualquer tempo, que será analisado conforme as regras vigentes na data do novo pedido, "
        "resguardado o direito adquirido (Emenda à Lei Orgânica nº 28/2025, art. 5º).\n"
        "Barra Mansa, {hoje_extenso}.\n"
        "Requerente: ______________________________________   Data: ____/____/______"),
    (DocumentKind.RELATORIO_INSTRUCAO, ""): (
        "RELATÓRIO DE INSTRUÇÃO – {tipo_concessao_maiusculo}",
        "Processo Administrativo nº {processo}. Interessado(a): {nome}, CPF {cpf}, matrícula {matricula}, "
        "cargo {cargo} ({orgao}).\n"
        "Enquadramento: {regra} – {fundamento}. Requisitos cumpridos em {data_cumprimento}.\n"
        "Tempo de contribuição apurado na data de validade: {tempo_contribuicao}.\n"
        "Data de validade: {data_validade}. Data-base da média e do tempo: {data_base}.\n"
        "Forma de cálculo: {forma_calculo}.\n"
        "Média dos 80% maiores salários: R$ {media80}; média de 100%: R$ {media100} ({competencias} competências); "
        "remuneração do cargo efetivo: R$ {ultima_remuneracao}.\n"
        "Valor apurado: R$ {valor} ({valor_extenso}), {forma_proventos}, {paridade}.\n"
        "{cotas}\n"
        "Encaminhe-se à Procuradoria Jurídica para parecer.\n"
        "Barra Mansa, {hoje_extenso}."),
    (DocumentKind.PARECER_JURIDICO, ""): (
        "PARECER JURÍDICO – PROCESSO Nº {processo}",
        "Assunto: {tipo_concessao} de {nome}, CPF {cpf}.\n"
        "{parecer}\n"
        "Barra Mansa, {hoje_extenso}."),
    (DocumentKind.PARECER_CONTROLE, ""): (
        "ANÁLISE DE CONFORMIDADE – PROCESSO Nº {processo}",
        "Assunto: {tipo_concessao} de {nome}, CPF {cpf}. Valor: R$ {valor}.\n"
        "{conformidade}\n"
        "Encaminhe-se à Presidência.\n"
        "Barra Mansa, {hoje_extenso}."),
}
# Quem assina cada documento padrão (os demais: Presidente)
DEFAULT_SIGNER = {DocumentKind.TERMO_DESISTENCIA: "Setor de Benefícios – {sigla}",
                  DocumentKind.RELATORIO_INSTRUCAO: "Setor de Benefícios – {sigla}",
                  DocumentKind.PARECER_JURIDICO: "Procuradoria Jurídica – {sigla}",
                  DocumentKind.PARECER_CONTROLE: "Controle Interno – {sigla}"}
# Minutas usam o mesmo texto da portaria correspondente (sem número e com a marca de MINUTA)
MINUTA_OF = {DocumentKind.MINUTA_PORTARIA: DocumentKind.PORTARIA,
             DocumentKind.MINUTA_ANULACAO: DocumentKind.PORTARIA_ANULACAO}


def get_template(institution, kind: str, concession_kind: str) -> tuple[str, str, str]:
    base_kind = MINUTA_OF.get(kind, kind)
    t = (DocumentTemplate.objects.filter(institution=institution, kind=base_kind, active=True,
                                         concession_kind=concession_kind).first()
         or DocumentTemplate.objects.filter(institution=institution, kind=base_kind, active=True,
                                            concession_kind="").first())
    if t:
        return t.title, t.body, t.signer
    title, body = (DEFAULT_TEMPLATES.get((base_kind, concession_kind)) or DEFAULT_TEMPLATES[(base_kind, "")])
    return title, body, DEFAULT_SIGNER.get(base_kind, "Presidente do {sigla}")


# ------------------------------------------------------------------ dados da concessão
def context(c, *, number: str = "", act_date=None, reason: str = "") -> dict:
    inst = c.institution
    p = c.person
    link = c.link
    benefit = c.revised_benefit
    ctx = {
        "instituicao": inst.name, "instituicao_maiusculo": inst.name.upper(), "sigla": inst.short_name,
        "processo": c.process.code if c.process_id else "", "numero_portaria": number or "_____/______",
        "data_portaria": br_date(act_date), "data_portaria_extenso": long_date(act_date).upper() if act_date else
        "____ DE __________ DE ______", "hoje_extenso": long_date(timezone.localdate()),
        "nome": p.name, "cpf": _cpf(p.cpf), "matricula": link.registration if link else "",
        "cargo": link.position.name if link and link.position_id else "", "orgao": str(link.organ) if link else "",
        "modalidade": (c.get_modality_display() or "").lower(), "modalidade_maiusculo": (c.get_modality_display() or "").upper(),
        "regra": c.rule.name if c.rule_id else "", "fundamento": c.legal_basis or "", "fundamento_local": c.local_legal_basis or "",
        "data_validade": br_date(c.validity_date), "valor": br_money(c.total or 0),
        "valor_extenso": reais_por_extenso(c.total or 0),
        "forma_proventos": "integralidade da remuneração do cargo efetivo" if c.integralidade else
        "base na média das remunerações de contribuição",
        "paridade": "com revisão pela paridade com os servidores em atividade" if c.paridade else
        "com reajuste para preservar o valor real (sem paridade)",
        "instituidor": p.name, "cpf_instituidor": _cpf(p.cpf), "data_obito": br_date(c.death_date),
        "tipo_concessao": c.get_kind_display().lower(), "motivo_indeferimento": reason or c.rejection_reason or "",
        "tce_processo": c.tce_reference or "_____", "motivo_revisao": c.revision_reason or "",
        "motivo_desistencia": c.withdrawal_reason or "",
        "tipo_concessao_maiusculo": c.get_kind_display().upper(),
        "parecer": c.opinion or "", "conformidade": c.compliance_opinion or "",
        "data_cumprimento": br_date(_iso(((c.analysis or {}).get("regra") or {}).get("data_cumprimento"))),
        "tempo_contribuicao": _tempo(c), "data_base": br_date(_iso((c.analysis or {}).get("data_base_calculo"))),
        "forma_calculo": (c.calculation or {}).get("forma") or (c.calculation or {}).get("base_forma") or "—",
        "media80": _money_txt((c.calculation or {}).get("media80")),
        "media100": _money_txt((c.calculation or {}).get("media100")),
        "competencias": (c.calculation or {}).get("competencias") or "—",
        "ultima_remuneracao": _money_txt((c.calculation or {}).get("ultima_remuneracao")),
        "portaria_original": (f"{c.get_act_nature_display()} nº {c.act_number}, de {br_date(c.act_date)}"
                              if c.act_number else "portaria de concessão"),
        "reimplantacao": (", devendo o(a) servidor(a) retornar à folha de pagamento dos servidores ativos do órgão de "
                          "origem" if c.kind == ConcessionKind.APOSENTADORIA else ""),
    }
    if c.kind == ConcessionKind.PENSAO:
        ctx["cotas"] = "\n".join(
            f"   {q.dependent.name}, CPF {_cpf(q.dependent.cpf)}, {q.get_relation_display().lower()} – cota de "
            f"{str(q.percent).rstrip('0').rstrip('.')}% (R$ {br_money(q.amount)}), {q.get_duration_display().lower()}"
            + (f" até {br_date(q.end_date)}" if q.end_date else "") + ";"
            for q in c.quotas.select_related("dependent"))
        ctx["nome"] = ", ".join(q.dependent.name for q in c.quotas.select_related("dependent")) or p.name
    if benefit is not None:
        ctx["nome"] = benefit.beneficiary.name
        ctx["cpf"] = _cpf(benefit.beneficiary.cpf)
        ctx["beneficio_revisto"] = benefit.get_kind_display().lower()
        ctx["valor_anterior"] = br_money(benefit.amount)
    ctx.setdefault("cotas", "")
    ctx.setdefault("beneficio_revisto", "")
    ctx.setdefault("valor_anterior", "")
    return ctx


def _iso(v):
    from datetime import date

    try:
        return date.fromisoformat(v) if v else None
    except (TypeError, ValueError):
        return None


def _money_txt(v) -> str:
    try:
        return br_money(v) if v not in (None, "") else "—"
    except Exception:  # valor não numérico na memória de cálculo
        return str(v)


def _tempo(c) -> str:
    t = ((c.analysis or {}).get("tempo") or {}).get("contribuicao") or {}
    days = t.get("dias") if isinstance(t, dict) else t
    if days in (None, ""):
        return "—"
    y, rest = divmod(int(days), 365)
    return f"{y} anos e {rest} dias ({int(days)} dias)"


def _cpf(v: str) -> str:
    d = re.sub(r"\D", "", v or "")
    return f"{d[:3]}.{d[3:6]}.{d[6:9]}-{d[9:]}" if len(d) == 11 else v


def _fill(text: str, ctx: dict) -> str:
    return re.sub(r"\{([a-z_]+)\}", lambda m: str(ctx.get(m.group(1), m.group(0))), text)


# ------------------------------------------------------------------ geração do .docx
def render_docx(title: str, body: str, signer: str, ctx: dict, draft: bool) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Pt, RGBColor

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    head = doc.add_paragraph()
    head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = head.add_run(ctx["instituicao_maiusculo"])
    r.bold = True
    if draft:
        m = doc.add_paragraph()
        m.alignment = WD_ALIGN_PARAGRAPH.CENTER
        mr = m.add_run("MINUTA — documento sem validade até a assinatura da Presidência")
        mr.bold = True
        mr.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    tr = t.add_run(_fill(title, ctx))
    tr.bold = True
    for line in _fill(body, ctx).split("\n"):
        para = doc.add_paragraph(line)
        para.alignment = WD_ALIGN_PARAGRAPH.CENTER if line.strip() == "RESOLVE:" else WD_ALIGN_PARAGRAPH.JUSTIFY
        para.paragraph_format.first_line_indent = Pt(0) if line.startswith("   ") else Pt(36)
    doc.add_paragraph()
    sig = doc.add_paragraph()
    sig.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sig.add_run("_______________________________________\n" + _fill(signer, ctx))
    foot = doc.add_paragraph()
    foot.alignment = WD_ALIGN_PARAGRAPH.LEFT
    fr = foot.add_run(f"Documento gerado pelo SIST FOPAG em {timezone.localtime():%d/%m/%Y %H:%M} — "
                      f"processo {ctx['processo']}.")
    fr.font.size = Pt(8)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def generate(c, kind: str, *, number: str = "", act_date=None, reason: str = "") -> ConcessionDocument:
    """Gera o documento, guarda com hash e anexa ao processo do Protocolo."""
    from apps.protocolo.models import Attachment

    title, body, signer = get_template(c.institution, kind, c.kind)
    ctx = context(c, number=number, act_date=act_date, reason=reason)
    content = render_docx(title, body, signer, ctx, draft=kind in MINUTA_OF)
    digest = hashlib.sha256(content).hexdigest()
    version = c.documents.filter(kind=kind).count() + 1
    label = DocumentKind(kind).label
    name = f"{label} - {c.process.code.replace('/', '-')} - v{version}.docx".replace(" ", "_")
    ctype = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    att = Attachment(institution=c.institution, process=c.process,
                     description=f"{label} (gerado pelo sistema, versão {version})", file_name=name,
                     file_size=len(content), content_type=ctype, sha256=digest)
    att.file.save(name, ContentFile(content), save=False)
    att.save()
    doc = ConcessionDocument(institution=c.institution, concession=c, kind=kind, version=version, file_name=name,
                             sha256=digest, attachment=att)
    doc.file.save(name, ContentFile(content), save=False)
    doc.save()
    return doc
