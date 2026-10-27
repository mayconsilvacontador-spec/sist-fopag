import csv
import io
from datetime import datetime, time

from django.http import HttpResponse
from django.utils import timezone
from django.utils.dateparse import parse_date
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from rest_framework import serializers, viewsets
from rest_framework.decorators import action
from rest_framework.response import Response

from apps.core.views import InstitutionScopedMixin

from .models import AuditAction, AuditLog
from .services import log_event

# Categorias do menu AUDITORIA -> filtros
CATEGORIES = {
    "alteracoes": {"action__in": [AuditAction.CREATE, AuditAction.UPDATE]},
    "calculos": {"action": AuditAction.CALCULATE},
    "exclusoes": {"action__in": [AuditAction.DELETE, AuditAction.CANCEL]},
    "usuarios": {"module__in": ["admin.usuarios", "admin.perfis", "admin.permissoes"]},
    "acessos": {"action__in": [AuditAction.LOGIN, AuditAction.LOGIN_FAILED, AuditAction.LOGOUT,
                               AuditAction.PASSWORD_CHANGE]},
    "fechamentos": {"action": AuditAction.CLOSE},
    "reaberturas": {"action": AuditAction.REOPEN},
    "exportacoes": {"action": AuditAction.EXPORT},
    "permissoes": {"action": AuditAction.PERMISSION_CHANGE},
}

EXPORT_LIMIT = 50000


class AuditLogSerializer(serializers.ModelSerializer):
    action_display = serializers.CharField(source="get_action_display", read_only=True)
    user_full_name = serializers.CharField(source="user.full_name", read_only=True, default="")

    class Meta:
        model = AuditLog
        fields = [
            "id", "timestamp", "username", "user_full_name", "ip", "user_agent", "request_id", "source",
            "action", "action_display", "module", "object_type", "object_id", "object_repr",
            "old_values", "new_values", "justification", "extra", "record_hash",
        ]


class AuditLogViewSet(InstitutionScopedMixin, viewsets.ReadOnlyModelViewSet):
    """
    Consulta da trilha de auditoria (somente leitura).

    Filtros: categoria, action, module, username, object_type, object_id,
    data_inicio, data_fim (dd/mm/aaaa ou aaaa-mm-dd), search.
    """

    module_code = "auditoria"
    serializer_class = AuditLogSerializer
    queryset = AuditLog.objects.select_related("user").all()
    search_fields = ["object_repr", "username", "object_id", "justification", "ip"]
    ordering_fields = ["timestamp", "action", "module", "username"]
    action_permissions = {"export": "exportar", "verify": "visualizar"}
    filterset_fields = []

    def get_queryset(self):
        from django.db.models import Q

        institution = self.get_institution()
        qs = AuditLog.objects.select_related("user").filter(
            Q(institution=institution) | Q(institution__isnull=True)
        )
        p = self.request.query_params
        cat = p.get("categoria")
        if cat in CATEGORIES:
            qs = qs.filter(**CATEGORIES[cat])
        for field in ("action", "module", "object_type", "object_id"):
            if p.get(field):
                qs = qs.filter(**{field: p[field]})
        if p.get("username"):
            qs = qs.filter(username__icontains=p["username"])
        start = _parse_br_date(p.get("data_inicio"))
        end = _parse_br_date(p.get("data_fim"))
        tz = timezone.get_current_timezone()
        if start:
            qs = qs.filter(timestamp__gte=timezone.make_aware(datetime.combine(start, time.min), tz))
        if end:
            qs = qs.filter(timestamp__lte=timezone.make_aware(datetime.combine(end, time.max), tz))
        return qs.order_by("-timestamp", "-id")

    @action(detail=True, methods=["get"])
    def verify(self, request, pk=None):
        log = self.get_object()
        ok = log.verify()
        return Response({
            "id": log.pk, "integro": ok,
            "detail": "Hash confere com o conteúdo gravado." if ok else
                      "ATENÇÃO: o conteúdo não confere com o hash registrado.",
        })

    @action(detail=False, methods=["get"])
    def export(self, request):
        fmt = request.query_params.get("formato", "csv").lower()
        qs = self.filter_queryset(self.get_queryset())
        total = qs.count()
        rows = list(qs[:EXPORT_LIMIT])
        log_event(
            AuditAction.EXPORT, module="auditoria", object_type="audit.auditlog",
            object_repr=f"Exportação da auditoria ({fmt.upper()})",
            extra={"filtros": dict(request.query_params.items()), "registros": len(rows), "total": total},
        )
        header = ["ID", "Data/Hora", "Usuário", "IP", "Ação", "Módulo", "Tipo", "ID do registro",
                  "Registro", "Valores anteriores", "Valores novos", "Justificativa", "Hash"]
        data = [
            [
                r.pk, timezone.localtime(r.timestamp).strftime("%d/%m/%Y %H:%M:%S"), r.username, r.ip or "",
                r.get_action_display(), r.module, r.object_type, r.object_id, r.object_repr,
                _json(r.old_values), _json(r.new_values), r.justification, r.record_hash,
            ]
            for r in rows
        ]
        stamp = timezone.localtime().strftime("%Y%m%d_%H%M%S")
        if fmt == "xlsx":
            return _xlsx_response(header, data, f"auditoria_{stamp}.xlsx")
        return _csv_response(header, data, f"auditoria_{stamp}.csv")


def _parse_br_date(value):
    if not value:
        return None
    if "/" in value:
        try:
            return datetime.strptime(value, "%d/%m/%Y").date()
        except ValueError:
            return None
    return parse_date(value)


def _json(value):
    if value in (None, {}, []):
        return ""
    import json
    return json.dumps(value, ensure_ascii=False, default=str)


def _safe(value):
    """Evita injeção de fórmulas em planilhas (CSV/Excel injection)."""
    if isinstance(value, str) and value[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + value
    return value


def _csv_response(header, data, filename):
    data = [[_safe(v) for v in row] for row in data]
    buffer = io.StringIO()
    writer = csv.writer(buffer, delimiter=";", quoting=csv.QUOTE_MINIMAL)
    writer.writerow(header)
    writer.writerows(data)
    # BOM UTF-8 para abrir corretamente no Excel
    resp = HttpResponse("﻿" + buffer.getvalue(), content_type="text/csv; charset=utf-8")
    resp["Content-Disposition"] = f'attachment; filename="{filename}"'
    return resp


def _xlsx_response(header, data, filename):
    data = [[_safe(v) for v in row] for row in data]
    wb = Workbook()
    ws = wb.active
    ws.title = "Auditoria"
    ws.append(header)
    fill = PatternFill("solid", fgColor="1F4E8C")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(vertical="center")
    for row in data:
        ws.append(row)
    widths = [8, 20, 18, 16, 18, 22, 22, 38, 40, 50, 50, 40, 66]
    for idx, w in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = w
    ws.freeze_panes = "A2"
    out = io.BytesIO()
    wb.save(out)
    resp = HttpResponse(
        out.getvalue(),
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    )
    resp["Content-Disposition"] = f'attachment; filename="{filename}"'
    return resp
