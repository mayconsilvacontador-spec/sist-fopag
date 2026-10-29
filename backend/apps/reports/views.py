"""
GET reports/?modulo=relatorios.folha     -> relatórios do módulo (só os que o usuário pode ver)
GET reports/<codigo>/?filtros...         -> colunas, linhas e totais
GET reports/<codigo>/export/?formato=xlsx|csv&filtros...  -> arquivo (exige permissão "exportar"; vai para a auditoria)
"""
import csv
import io
from datetime import date
from decimal import Decimal

from django.http import Http404, HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from rest_framework.response import Response
from rest_framework.views import APIView

from apps.audit.services import log_event

from .registry import REPORTS, jsonable


def _get(code):
    rep = REPORTS.get(code)
    if rep is None:
        raise Http404
    return rep


class _ReportView(APIView):
    """A permissão é a do módulo do relatório (relatorios.*), conferida pela ModulePermission."""

    def initial(self, request, *args, **kwargs):
        code = kwargs.get("code")
        rep = _get(code) if code else None
        # ?modulo= indica a tela de onde o relatório foi aberto (tem de ser um dos módulos do relatório)
        wanted = request.query_params.get("modulo")
        self.module_code = (wanted if wanted in rep.modules else rep.module) if rep else None
        super().initial(request, *args, **kwargs)


class ReportListView(APIView):
    def get(self, request):
        modulo = request.query_params.get("modulo", "")
        keys = getattr(request, "permission_keys", set())
        if modulo:
            out = [{**r.meta(), "modulo": modulo} for r in REPORTS.values()
                   if modulo in r.modules and f"{modulo}:visualizar" in keys]
        else:
            out = [r.meta() for r in REPORTS.values() if f"{r.module}:visualizar" in keys]
        return Response(out)


def _run(rep, request):
    params = {k: v for k, v in request.query_params.items() if k not in ("formato", "modulo")}
    data = rep.run(request.institution, params)
    data["linhas"] = [{k: jsonable(v) for k, v in row.items()} for row in data["linhas"]]
    data["totais"] = {k: jsonable(v) for k, v in data["totais"].items()}
    return {**rep.meta(), **data, "parametros": params, "instituicao": request.institution.name,
            "gerado_em": timezone.localtime().isoformat()}


class ReportRunView(_ReportView):
    def get(self, request, code):
        return Response(_run(_get(code), request))


class ReportExportView(_ReportView):
    action_permissions = {"GET": "exportar"}

    def get(self, request, code):
        rep = _get(code)
        data = _run(rep, request)
        fmt = request.query_params.get("formato", "xlsx")
        stamp = timezone.localtime().strftime("%Y%m%d-%H%M")
        name = f"{rep.code}-{stamp}"
        log_event("EXPORT", module=self.module_code, object_repr=f"Relatório {rep.title} ({fmt.upper()})",
                  extra={"filtros": data["parametros"], "linhas": len(data["linhas"])})
        if fmt == "csv":
            return _csv(data, name)
        return _xlsx(data, name)


def _cell(v, kind):
    if v is None or v == "":
        return ""
    if kind == "money" or kind == "percent":
        return Decimal(str(v))
    if kind == "date":
        return date.fromisoformat(v)
    if kind == "bool":
        return "Sim" if v else "Não"
    return v


def _xlsx(data, name):
    wb = Workbook()
    ws = wb.active
    ws.title = "Relatório"
    cols = data["colunas"]
    ws.append([data["instituicao"]])
    ws.append([data["titulo"]])
    ws["A2"].font = Font(bold=True, size=13)
    filtros = ", ".join(f"{k}: {v}" for k, v in data["parametros"].items() if v) or "sem filtros"
    ws.append([f"Filtros: {filtros} · gerado em {timezone.localtime():%d/%m/%Y %H:%M}"])
    ws.append([])
    ws.append([c["label"] for c in cols])
    head = ws.max_row
    for c in ws[head]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E6E")
        c.alignment = Alignment(vertical="center", wrap_text=True)
    fmt = {"money": '#,##0.00', "percent": '0.0', "date": "DD/MM/YYYY", "int": "0"}
    for row in data["linhas"]:
        ws.append([_cell(row.get(c["key"]), c["tipo"]) for c in cols])
        for i, c in enumerate(cols, 1):
            if c["tipo"] in fmt:
                ws.cell(ws.max_row, i).number_format = fmt[c["tipo"]]
            if row.get("_subtotal"):  # subtotal (ex.: por órgão) em negrito
                ws.cell(ws.max_row, i).font = Font(bold=True)
    if data["totais"]:
        ws.append([_cell(data["totais"].get(c["key"]), c["tipo"]) if c["key"] in data["totais"] else ""
                   for c in cols])
        for i, c in enumerate(cols, 1):
            cell = ws.cell(ws.max_row, i)
            cell.font = Font(bold=True)
            if c["tipo"] in fmt and not isinstance(cell.value, str):
                cell.number_format = fmt[c["tipo"]]
    for note in data.get("observacoes") or []:
        ws.append([note])
    for i, c in enumerate(cols, 1):
        width = max([len(c["label"])] + [len(str(r.get(c["key"]) or "")) for r in data["linhas"][:300]]) + 2
        ws.column_dimensions[ws.cell(head, i).column_letter].width = min(max(width, 8), 60)
    ws.freeze_panes = ws.cell(head + 1, 1)
    buf = io.BytesIO()
    wb.save(buf)
    resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    resp["Content-Disposition"] = f'attachment; filename="{name}.xlsx"'
    return resp


def _csv(data, name):
    """CSV para Excel em português: separador ';' e vírgula decimal."""
    buf = io.StringIO()
    w = csv.writer(buf, delimiter=";")
    cols = data["colunas"]
    w.writerow([c["label"] for c in cols])

    def fmt(v, kind):
        v = _cell(v, kind)
        if isinstance(v, Decimal):
            return f"{v:.2f}".replace(".", ",")
        if isinstance(v, date):
            return v.strftime("%d/%m/%Y")
        return v

    for row in data["linhas"]:
        w.writerow([fmt(row.get(c["key"]), c["tipo"]) for c in cols])
    if data["totais"]:
        w.writerow([fmt(data["totais"].get(c["key"]), c["tipo"]) if c["key"] in data["totais"] else "" for c in cols])
    resp = HttpResponse("﻿" + buf.getvalue(), content_type="text/csv; charset=utf-8")
    resp["Content-Disposition"] = f'attachment; filename="{name}.csv"'
    return resp
