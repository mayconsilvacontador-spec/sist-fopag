from django.db.models import Case, CharField, Exists, OuterRef, Q, Value, When
from django.http import FileResponse
from rest_framework import mixins, serializers, viewsets
from rest_framework.decorators import action
from rest_framework.parsers import FormParser, JSONParser, MultiPartParser
from rest_framework.response import Response

from apps.audit.services import log_event
from apps.core.exceptions import BusinessRuleError
from apps.core.validators import only_digits
from apps.core.views import InstitutionScopedMixin

from . import services
from .models import Address, BankAccount, Dependency, Document, DocumentType, ImportBatch, ImportStatus, Person
from .serializers import (
    AddressSerializer, BankAccountSerializer, DependencySerializer, DocumentSerializer, PersonListSerializer,
    PersonSerializer,
)


class PersonViewSet(InstitutionScopedMixin, mixins.ListModelMixin, mixins.RetrieveModelMixin,
                    mixins.CreateModelMixin, mixins.UpdateModelMixin, viewsets.GenericViewSet):
    """
    Cadastro Único. Sem exclusão. Filtros: ?search= (nome, CPF, matrícula),
    ?tipo=servidor|aposentado|dependente|falecido.
    """

    module_code = "cadastros.segurados"
    queryset = Person.objects.all()
    ordering_fields = ["name", "birth_date", "created_at"]

    def get_serializer_class(self):
        return PersonListSerializer if self.action == "list" else PersonSerializer

    def get_queryset(self):
        from apps.employment.models import ACTIVE_STATUSES, EmploymentLink, LinkStatus

        qs = super().get_queryset()
        links = EmploymentLink.objects.filter(person=OuterRef("pk"))
        qs = qs.annotate(
            has_active=Exists(links.filter(status__in=ACTIVE_STATUSES)),
            has_retired=Exists(links.filter(status=LinkStatus.APOSENTADO)),
            is_dependent=Exists(Dependency.objects.filter(dependent=OuterRef("pk"), end_date__isnull=True)),
        ).annotate(situacao_label=Case(
            When(has_active=True, then=Value("Servidor ativo")),
            When(has_retired=True, then=Value("Aposentado(a)")),
            When(is_dependent=True, then=Value("Dependente")),
            default=Value(""), output_field=CharField(),
        ))
        term = (self.request.query_params.get("search") or "").strip()
        if term:
            digits = only_digits(term)
            cond = Q(name__icontains=term) | Q(social_name__icontains=term)
            if len(digits) >= 3:
                cond |= Q(cpf__startswith=digits) | Q(employment_links__registration__icontains=digits)
            qs = qs.filter(cond).distinct()
        tipo = self.request.query_params.get("tipo")
        if tipo == "servidor":
            qs = qs.filter(has_active=True)
        elif tipo == "aposentado":
            qs = qs.filter(has_retired=True)
        elif tipo == "dependente":
            qs = qs.filter(is_dependent=True)
        elif tipo == "falecido":
            qs = qs.filter(death_date__isnull=False)
        if self.request.query_params.get("sem_nome") == "1":
            qs = qs.filter(name="")
        return qs

    action_permissions = {"receita": "incluir", "fill_from_receita": "editar", "admissao": "incluir",
                          "cronologia": "visualizar", "ficha": "visualizar", "pasta": "visualizar", "foto": "visualizar"}

    @action(detail=False, methods=["post"])
    def admissao(self, request):
        """Admissão: pessoa + vínculo + nomeação/posse + endereço + conta, de uma vez."""
        from .lifecycle import admit

        return Response(admit(self.get_institution(), request.data, request.user), status=201)

    @action(detail=True, methods=["get"])
    def cronologia(self, request, pk=None):
        from .lifecycle import timeline

        return Response(timeline(self.get_object()))

    @action(detail=True, methods=["get"])
    def pasta(self, request, pk=None):
        from .lifecycle import folder_extras

        return Response(folder_extras(self.get_object(), set(getattr(request, "permission_keys", set()))))

    @action(detail=True, methods=["get"])
    def ficha(self, request, pk=None):
        from .lifecycle import summary

        return Response(summary(self.get_object()))

    @action(detail=True, methods=["get", "post", "delete"], parser_classes=[MultiPartParser, FormParser, JSONParser])
    def foto(self, request, pk=None):
        """GET devolve a foto; POST (multipart: file) troca; DELETE remove. PNG/JPG/WEBP conferidos pelo conteúdo."""
        from django.http import Http404, HttpResponse

        from apps.institutions.branding import sniff

        person = self.get_object()
        if request.method == "GET":
            if not person.photo:
                raise Http404
            with person.photo.open("rb") as f:
                return HttpResponse(f.read(), content_type=person.photo_type or "image/jpeg")
        if "cadastros.segurados:editar" not in getattr(request, "permission_keys", set()):
            return Response({"detail": "Sem permissão para alterar a foto."}, status=403)
        if request.method == "DELETE":
            person.photo.delete(save=False)
            person.photo, person.photo_type = "", ""
            person.save(update_fields=["photo", "photo_type", "updated_at"])
            log_event("UPDATE", module="cadastros.segurados", obj=person, object_repr=f"{person.name} — foto removida")
            return Response(status=204)
        upload = request.FILES.get("file")
        if upload is None:
            raise BusinessRuleError("Selecione a foto.")
        if upload.size > 5 * 1024 * 1024:
            raise BusinessRuleError("A foto deve ter no máximo 5 MB.")
        content = upload.read()
        kind = sniff(content)
        if kind is None:
            raise BusinessRuleError("A foto deve ser PNG, JPG ou WEBP.")
        from django.core.files.base import ContentFile

        if person.photo:
            person.photo.delete(save=False)
        person.photo.save(f"{person.pk}.{kind[1]}", ContentFile(content), save=False)
        person.photo_type = kind[0]
        person.save(update_fields=["photo", "photo_type", "updated_at"])
        log_event("UPDATE", module="cadastros.segurados", obj=person, object_repr=f"{person.name} — foto atualizada")
        return Response({"ok": True})

    @action(detail=False, methods=["get"], url_path="by-cpf")
    def by_cpf(self, request):
        """Busca exata por CPF — usada para reaproveitar o cadastro (evita duplicidade)."""
        cpf = only_digits(request.query_params.get("cpf"))
        person = self.get_queryset().filter(cpf=cpf).first()
        return Response(PersonListSerializer(person).data if person else None)

    @action(detail=False, methods=["get"])
    def receita(self, request):
        """Consulta um CPF na Receita Federal (Conecta gov.br) para pré-preencher o cadastro."""
        from . import receita

        cpf = only_digits(request.query_params.get("cpf"))
        data = receita.consult([cpf], request.user.cpf or "").get(cpf)
        log_event("SYSTEM", module="cadastros.segurados", object_type="receita.cpf",
                  object_repr="Consulta de CPF na Receita Federal (Conecta gov.br)",
                  extra={"cpf": f"***.{cpf[3:6]}.{cpf[6:9]}-**", "encontrado": bool(data)})
        if data is None:
            raise BusinessRuleError("CPF não encontrado na Receita Federal.")
        return Response(data)

    @action(detail=False, methods=["post"], url_path="fill-from-receita")
    def fill_from_receita(self, request):
        """Completa pela Receita o nome (e nome da mãe) das pessoas sem nome — até 50 por vez."""
        from . import receita

        people = list(Person.objects.filter(institution=self.get_institution(), name="")[:receita.MAX_BATCH])
        found = receita.consult([p.cpf for p in people], request.user.cpf or "")
        filled = 0
        for p in people:
            d = found.get(p.cpf)
            if d and d["name"]:
                p.name = d["name"][:150]
                p.social_name = p.social_name or d["social_name"][:150]
                p.mother_name = p.mother_name or d["mother_name"][:150]
                p.save()
                filled += 1
        log_event("SYSTEM", module="cadastros.segurados", object_type="receita.cpf",
                  object_repr="Preenchimento de nomes pela Receita Federal",
                  extra={"consultados": len(people), "preenchidos": filled})
        remaining = Person.objects.filter(institution=self.get_institution(), name="").count()
        return Response({"consultados": len(people), "preenchidos": filled, "restantes": remaining})


class PersonChildViewSet(InstitutionScopedMixin, mixins.ListModelMixin, mixins.RetrieveModelMixin,
                         mixins.CreateModelMixin, viewsets.GenericViewSet):
    """Base dos registros da pessoa: ?person=<id> obrigatório na listagem."""

    module_code = "cadastros.segurados"
    filterset_fields = ["person"]
    pagination_class = None

    def get_queryset(self):
        qs = super().get_queryset()
        if self.action == "list" and not self.request.query_params.get("person"):
            return qs.none()
        return qs

    def _person(self, serializer):
        person = serializer.validated_data["person"]
        if person.institution_id != self.get_institution().pk:
            raise serializers.ValidationError({"person": "Pessoa de outra instituição."})
        return person


class AddressViewSet(PersonChildViewSet):
    """Novo endereço encerra a vigência do anterior. Não há alteração nem exclusão."""

    serializer_class = AddressSerializer
    queryset = Address.objects.all()
    action_permissions = {"create": "editar"}

    def perform_create(self, serializer):
        data = dict(serializer.validated_data)
        person = self._person(serializer)
        data.pop("person")
        serializer.instance = services.add_address(person, **data)


class BankAccountViewSet(PersonChildViewSet):
    serializer_class = BankAccountSerializer
    queryset = BankAccount.objects.select_related("bank").all()
    action_permissions = {"create": "editar"}

    def perform_create(self, serializer):
        data = dict(serializer.validated_data)
        person = self._person(serializer)
        data.pop("person")
        serializer.instance = services.add_bank_account(person, **data)


class DependencyViewSet(InstitutionScopedMixin, mixins.ListModelMixin, mixins.RetrieveModelMixin,
                        mixins.CreateModelMixin, mixins.UpdateModelMixin, viewsets.GenericViewSet):
    """Dependências. ?holder=<id> (dependentes do titular) ou ?dependent=<id>. Encerrar em vez de excluir."""

    module_code = "cadastros.dependentes"
    serializer_class = DependencySerializer
    queryset = Dependency.objects.select_related("holder", "dependent").all()
    filterset_fields = ["holder", "dependent"]

    def perform_create(self, serializer):
        d = serializer.validated_data
        inst = self.get_institution()
        if d["holder"].institution_id != inst.pk or d["dependent"].institution_id != inst.pk:
            raise serializers.ValidationError("Pessoa de outra instituição.")
        services.validate_dependency(d["holder"], d["dependent"], d["kind"], d["start_date"])
        serializer.save(institution=inst)

    def perform_update(self, serializer):
        if {"holder", "dependent"} & set(serializer.validated_data):
            raise BusinessRuleError("Titular e dependente não podem ser alterados. Encerre e crie nova dependência.")
        serializer.save()


class ImportBatchViewSet(InstitutionScopedMixin, mixins.ListModelMixin, mixins.RetrieveModelMixin,
                         viewsets.GenericViewSet):
    """
    Importação de segurados (leiaute SPREV).
    POST (multipart: files[], source_system) -> valida sem gravar e devolve o relatório;
    POST {id}/confirm/ -> importa; POST {id}/discard/ -> descarta.
    """

    module_code = "cadastros.importacao"
    queryset = ImportBatch.objects.select_related("created_by").all()
    parser_classes = [MultiPartParser, FormParser, JSONParser]
    action_permissions = {"confirm": "incluir", "discard": "incluir", "modelo_servidores": "visualizar"}
    MAX_FILE_BYTES = 50 * 1024 * 1024

    @action(detail=False, methods=["get"], url_path="modelo-servidores")
    def modelo_servidores(self, request):
        """Tabela padrão de importação de servidores (Excel): aba SERVIDORES (só o cabeçalho) + aba INSTRUÇÕES."""
        import io

        import openpyxl
        from django.http import HttpResponse
        from openpyxl.styles import Font, PatternFill

        from .importers.complementar import SERVIDORES_COLUMNS

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "SERVIDORES"
        ws.append([c for c, _ in SERVIDORES_COLUMNS])
        required = {"CPF", "NOME", "DATA_NASCIMENTO", "SEXO", "ORGAO", "CARGO", "MATRICULA", "DATA_NOMEACAO"}
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="C2410C" if cell.value in required else "1E3A5F")
            ws.column_dimensions[cell.column_letter].width = max(14, len(str(cell.value)) + 4)
            cell.number_format = "@"
        ws.freeze_panes = "A2"
        for col in range(1, len(SERVIDORES_COLUMNS) + 1):  # colunas como texto (CPF/matrícula sem perder zeros)
            for row in range(2, 2001):
                ws.cell(row=row, column=col).number_format = "@"
        ins = wb.create_sheet("INSTRUCOES")
        ins.append(["Coluna", "Como preencher"])
        for cell in ins[1]:
            cell.font = Font(bold=True)
        for c, help_text in SERVIDORES_COLUMNS:
            ins.append([c, ("OBRIGATÓRIA. " if c in required and not help_text.startswith("Obrigat") else "") + help_text])
        ins.append([])
        ins.append(["Dicas", "Uma linha por vínculo (servidor com dois cargos = duas linhas com o mesmo CPF). "
                             "Colunas laranja são obrigatórias. Salve como .xlsx ou .csv (separado por ;) e envie em "
                             "Servidores › Importar. O sistema valida tudo antes de gravar; nada é sobrescrito."])
        ins.column_dimensions["A"].width = 28
        ins.column_dimensions["B"].width = 110
        buf = io.BytesIO()
        wb.save(buf)
        resp = HttpResponse(buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        resp["Content-Disposition"] = 'attachment; filename="SIST_tabela_padrao_servidores.xlsx"'
        return resp

    def get_serializer_class(self):
        from .serializers import ImportBatchSerializer
        return ImportBatchSerializer

    def _tables(self, batch):
        from django.core.files.storage import default_storage

        from .importers.sprev import read_table

        tables = []
        for f in batch.files:
            with default_storage.open(f["caminho"], "rb") as fh:
                tables.append((f["nome"], read_table(fh.read(), f["nome"])))
        return tables

    def create(self, request, *args, **kwargs):
        import hashlib

        from django.core.files.base import ContentFile
        from django.core.files.storage import default_storage

        from .importers.sprev import SprevImporter, read_table

        uploads = request.FILES.getlist("files")
        source = (request.data.get("source_system") or "").strip()
        if not uploads:
            raise serializers.ValidationError({"files": "Selecione ao menos um arquivo."})
        if not source:
            raise serializers.ValidationError({"source_system": "Informe o sistema ou órgão de origem."})
        for up in uploads:
            if not up.name.lower().endswith((".csv", ".txt", ".xlsx")):
                raise BusinessRuleError(f"{up.name}: envie arquivos CSV ou Excel (.xlsx).")
            if up.size > self.MAX_FILE_BYTES:
                raise BusinessRuleError(f"{up.name}: arquivo maior que 50 MB.")
        institution = self.get_institution()
        batch = ImportBatch.objects.create(institution=institution, source_system=source[:150])
        files, tables = [], []
        for up in uploads:
            content = up.read()
            sha = hashlib.sha256(content).hexdigest()
            path = default_storage.save(f"importacoes/{batch.pk}/{sha[:12]}_{up.name[-80:]}", ContentFile(content))
            rows = read_table(content, up.name)
            files.append({"nome": up.name, "caminho": path, "sha256": sha, "linhas": len(rows)})
            tables.append((up.name, rows))
        batch.files = files
        batch.report = SprevImporter(institution, batch, source).run(tables, apply=False)
        batch.save()
        return Response(self.get_serializer(batch).data, status=201)

    @action(detail=True, methods=["post"])
    def confirm(self, request, pk=None):
        from django.utils import timezone

        from .importers.sprev import SprevImporter

        batch = self.get_object()
        if batch.status != ImportStatus.VALIDADO:
            raise BusinessRuleError("Este lote não está aguardando confirmação.")
        batch.result = SprevImporter(batch.institution, batch, batch.source_system).run(self._tables(batch), apply=True)
        batch.status, batch.imported_at = ImportStatus.IMPORTADO, timezone.now()
        batch.save()
        log_event("CREATE", module="cadastros.importacao", obj=batch,
                  object_repr=f"Importação de segurados — {batch.source_system}", extra=batch.result.get("resumo"))
        return Response(self.get_serializer(batch).data)

    @action(detail=True, methods=["post"])
    def discard(self, request, pk=None):
        batch = self.get_object()
        if batch.status != ImportStatus.VALIDADO:
            raise BusinessRuleError("Somente lotes não importados podem ser descartados.")
        batch.status = ImportStatus.DESCARTADO
        batch.save()
        return Response(self.get_serializer(batch).data)


class DocumentViewSet(PersonChildViewSet):
    """Upload (multipart: person, doc_type, file, description) e download de documentos."""

    serializer_class = DocumentSerializer
    queryset = Document.objects.select_related("created_by").all()
    parser_classes = [MultiPartParser, FormParser]
    action_permissions = {"create": "editar", "download": "visualizar"}

    def create(self, request, *args, **kwargs):
        person = Person.objects.filter(institution=self.get_institution(), pk=request.data.get("person")).first()
        if person is None:
            return Response({"person": ["Pessoa não encontrada."]}, status=400)
        doc_type = request.data.get("doc_type")
        if doc_type not in DocumentType.values:
            return Response({"doc_type": ["Tipo de documento inválido."]}, status=400)
        upload = request.FILES.get("file")
        if upload is None:
            return Response({"file": ["Selecione o arquivo."]}, status=400)
        doc = services.upload_document(person, doc_type, upload, request.data.get("description", ""))
        return Response(DocumentSerializer(doc).data, status=201)

    @action(detail=True, methods=["get"])
    def download(self, request, pk=None):
        doc = self.get_object()
        log_event("EXPORT", module="cadastros.segurados", obj=doc, object_repr=f"Download {doc.file_name}")
        return FileResponse(doc.file.open("rb"), as_attachment=True, filename=doc.file_name,
                            content_type=doc.content_type)
