"""Salvaguardas da concessão: segregação de funções, assinatura com senha, hash dos valores, documentos, TCE-RJ."""
import io
from datetime import date
from decimal import Decimal

import pytest
from django.utils import timezone

from apps.alerts.models import Alert
from apps.cadastros.models import Ente, Organ, Position
from apps.concessao.models import Benefit, Concession, ConcessionStep
from apps.employment.models import EmploymentLink
from apps.institutions.services import ensure_payroll_groups, set_parameter
from apps.people.models import Person
from apps.previdencia.rules import seed_rules
from apps.protocolo.models import Attachment, ProcessType, Sector
from tests.conftest import PASSWORD, anexar_checklist, enviar, historico

REFORM = date(2025, 6, 15)
ACK = {"acknowledge_warnings": True}
TEXTO = "Requisitos constitucionais e legais atendidos; fundamentação e cálculo corretos."


@pytest.fixture
def caso(institution, master, setores, client_for):
    seed_rules(institution, REFORM)
    set_parameter(institution, "previdencia.vigencia_reforma_local", REFORM.isoformat(), REFORM)
    set_parameter(institution, "rgps.teto", "8157.41", date(2025, 1, 1))
    ensure_payroll_groups(institution)
    organ = Organ.objects.create(institution=institution, ente=Ente.objects.create(institution=institution, name="M"),
                                 name="Prefeitura")
    pos = Position.objects.create(institution=institution, name="Agente", cargo_type=7)
    sector = Sector.objects.create(institution=institution, name="Benefícios", code="BEN")
    sector.members.add(master, setores["beneficios"])
    ptype = ProcessType.objects.create(institution=institution, name="Aposentadoria", deadline_days=60)
    p = Person.objects.create(institution=institution, cpf="52998224725", name="José Servidor",
                              birth_date=date(1958, 3, 10), sex=2)
    EmploymentLink.objects.create(institution=institution, person=p, organ=organ, position=pos, registration="10",
                                  public_service_start=date(1988, 1, 4), ente_start=date(1988, 1, 4),
                                  position_start=date(1988, 1, 4), career_start=date(1988, 1, 4),
                                  segment="FINANCEIRO", base_salary=Decimal("5000.00"))
    historico(institution, p, date(1994, 7, 1), date(2026, 8, 1), "5200.00")
    ben = client_for(setores["beneficios"])
    pid = ben.post("/api/v1/protocol/processes/", {"process_type": str(ptype.pk), "sector": str(sector.pk),
                                                   "subject": "Aposentadoria", "interested": str(p.pk),
                                                   "interested_name": "x"}, format="json").json()["id"]
    cid = ben.post("/api/v1/concessoes/", {"kind": "APOSENTADORIA", "person": str(p.pk), "process": pid},
                   format="json").json()["id"]
    assert ben.post(f"/api/v1/concessoes/{cid}/frame/", {"rule": "EC41_ART6", "validity_date": "2026-09-01",
                                                         "validity_class": 3}, format="json").status_code == 200
    assert ben.post(f"/api/v1/concessoes/{cid}/calculate/").status_code == 200
    return {"cid": cid, "person": p}


def test_fluxo_seguro_da_concessao(institution, master, setores, client_for, caso):
    cid = caso["cid"]
    c = {k: client_for(u) for k, u in setores.items()}
    today = timezone.localdate().isoformat()

    # Não existe mais "aprovar e implantar" direto: implantação antes da assinatura é recusada
    assert c["folha"].post(f"/api/v1/concessoes/{cid}/implant/", {"password": PASSWORD}, format="json").status_code == 409
    assert enviar(c["beneficios"], cid).status_code == 200

    # Cada ato é do setor competente: Benefícios não dá parecer jurídico nem assina
    assert c["beneficios"].post(f"/api/v1/concessoes/{cid}/legal/", {"favorable": True, "text": TEXTO,
                                "password": PASSWORD}, format="json").status_code == 403
    assert c["controle"].post(f"/api/v1/concessoes/{cid}/legal/", {"favorable": True, "text": TEXTO,
                              "password": PASSWORD}, format="json").status_code == 403
    # Senha errada: nada é registrado
    r = c["juridico"].post(f"/api/v1/concessoes/{cid}/legal/", {"favorable": True, "text": TEXTO,
                           "password": "errada"}, format="json")
    assert r.status_code == 409 and "Senha incorreta" in r.json()["detail"]
    assert Concession.objects.get(pk=cid).stage == "ANALISE_JURIDICA"

    # Devolução pelo Jurídico volta para Benefícios corrigir
    r = c["juridico"].post(f"/api/v1/concessoes/{cid}/legal/", {"favorable": False, "password": PASSWORD,
                           "text": "Juntar certidão de tempo de contribuição do INSS do período de 1980 a 1987."},
                           format="json")
    assert r.json()["stage"] == "CALCULADO" and r.json()["returned_count"] == 1
    assert Alert.objects.filter(institution=institution, code="CONCESSAO_PENDENTE", title__startswith="Processo devolvido",
                                status="ABERTO").exists()
    enviar(c["beneficios"], cid)
    r = c["juridico"].post(f"/api/v1/concessoes/{cid}/legal/", {"favorable": True, "text": TEXTO,
                           "password": PASSWORD}, format="json")
    assert r.json()["stage"] == "CONTROLE_INTERNO"

    # Minuta gerada automaticamente, anexada ao Protocolo e avisada a Benefícios
    docs = c["beneficios"].get(f"/api/v1/concessoes/{cid}/documents/").json()
    assert docs[0]["kind"] == "MINUTA_PORTARIA" and len(docs[0]["sha256"]) == 64
    assert Attachment.objects.filter(process__concessions__pk=cid, description__startswith="Minuta").exists()
    dl = c["beneficios"].get(f"/api/v1/concessoes/{cid}/documents/{docs[0]['id']}/download/")
    assert dl.status_code == 200 and dl.content[:2] == b"PK"  # .docx
    from docx import Document

    text = "\n".join(p.text for p in Document(io.BytesIO(dl.content)).paragraphs)
    assert "MINUTA" in text and "JOSÉ SERVIDOR" not in text and "José Servidor" in text
    assert "R$ 5.200,00 (cinco mil e duzentos reais)" in text
    assert Alert.objects.filter(institution=institution, code="CONCESSAO_PENDENTE", module="concessao.controle_interno",
                                status="ABERTO").exists()

    # Controle Interno confere; Presidência assina — a portaria numerada é gerada e anexada
    r = c["controle"].post(f"/api/v1/concessoes/{cid}/compliance/", {"conform": True, "password": PASSWORD,
                           "text": "Regras e memória de cálculo conferidas."}, format="json")
    assert r.json()["stage"] == "PRESIDENCIA"
    r = c["presidente"].post(f"/api/v1/concessoes/{cid}/sign/", {"nature": "PORTARIA", "number": "77/2026",
                             "act_date": today, "password": PASSWORD}, format="json")
    assert r.json()["stage"] == "ATO_EMITIDO" and len(r.json()["signed_hash"]) == 64
    port = next(d for d in c["beneficios"].get(f"/api/v1/concessoes/{cid}/documents/").json() if d["kind"] == "PORTARIA")
    content = c["beneficios"].get(f"/api/v1/concessoes/{cid}/documents/{port['id']}/download/").content
    text = "\n".join(p.text for p in Document(io.BytesIO(content)).paragraphs)
    assert "Nº 77/2026" in text and "MINUTA" not in text


def test_segregacao_de_funcoes(institution, master, setores, client_for, caso):
    cid = caso["cid"]
    m = client_for(master)  # MASTER tem todas as permissões
    c = {k: client_for(u) for k, u in setores.items()}
    enviar(c["beneficios"], cid)
    # O MASTER dá o parecer...
    r = m.post(f"/api/v1/concessoes/{cid}/legal/", {"favorable": True, "text": TEXTO, "password": PASSWORD}, format="json")
    assert r.status_code == 200
    # ...e por isso não pode também conferir nem assinar o mesmo processo
    r = m.post(f"/api/v1/concessoes/{cid}/compliance/", {"conform": True, "text": "Conferido sem ressalvas.",
               "password": PASSWORD}, format="json")
    assert r.status_code == 409 and "Segregação" in r.json()["detail"]
    c["controle"].post(f"/api/v1/concessoes/{cid}/compliance/", {"conform": True, "text": "Conferido sem ressalvas.",
                       "password": PASSWORD}, format="json")
    r = m.post(f"/api/v1/concessoes/{cid}/sign/", {"number": "1/2026", "act_date": timezone.localdate().isoformat(),
               "password": PASSWORD}, format="json")
    assert r.status_code == 409 and "Segregação" in r.json()["detail"]
    # Quem instruiu (Benefícios) não implanta
    today = timezone.localdate().isoformat()
    c["presidente"].post(f"/api/v1/concessoes/{cid}/sign/", {"number": "1/2026", "act_date": today,
                         "password": PASSWORD}, format="json")
    anexar_checklist(c["folha"], cid)
    c["folha"].post(f"/api/v1/concessoes/{cid}/publish/", {"publication_date": today, "media": "Notícia Oficial 10", "password": PASSWORD},
                    format="json")
    steps = {s["action"]: s for s in c["folha"].get(f"/api/v1/concessoes/{cid}/steps/").json()}
    assert steps["ASSINATURA"]["signed"] and len(steps["ASSINATURA"]["signature_hash"]) == 64
    assert {d["kind"] for d in c["folha"].get(f"/api/v1/concessoes/{cid}/documents/").json()} >= {"RELATORIO_INSTRUCAO", "PARECER_JURIDICO", "MINUTA_PORTARIA", "PARECER_CONTROLE", "PORTARIA"}

    # Valores alterados depois da assinatura: implantação bloqueada
    Concession.objects.filter(pk=cid).update(total=Decimal("9999.00"))
    r = c["folha"].post(f"/api/v1/concessoes/{cid}/implant/", {"password": PASSWORD}, format="json")
    assert r.status_code == 409 and "não conferem" in r.json()["detail"]
    Concession.objects.filter(pk=cid).update(total=Decimal("5200.00"))
    r = c["folha"].post(f"/api/v1/concessoes/{cid}/implant/", {"password": PASSWORD}, format="json")
    assert r.status_code == 200 and r.json()["stage"] == "IMPLANTADO"
    ben = Benefit.objects.get(concession_id=cid)
    assert ben.amount == Decimal("5200.00")
    assert ConcessionStep.objects.filter(concession_id=cid, action="IMPLANTACAO", signed=True).exists()

    # TCE-RJ: julgamento ILEGAL gera minuta de anulação; a Presidência anula com senha
    c["presidente"].post(f"/api/v1/concessoes/{cid}/tce-send/", {"sent_date": today, "reference": "TCE 123/2026"},
                         format="json")
    r = c["presidente"].post(f"/api/v1/concessoes/{cid}/tce-judgment/", {"result": "ILEGAL", "result_date": today,
                             "notes": "Tempo especial sem comprovação."}, format="json")
    assert r.status_code == 200
    assert "MINUTA_ANULACAO" in {d["kind"] for d in c["presidente"].get(f"/api/v1/concessoes/{cid}/documents/").json()}
    r = c["presidente"].post(f"/api/v1/concessoes/{cid}/annul/", {"number": "50/2026", "act_date": today,
                             "password": PASSWORD}, format="json")
    assert r.status_code == 200 and r.json()["stage"] == "ANULADO"
    ben.refresh_from_db()
    assert ben.status == "CESSADO"
    assert Concession.objects.get(pk=cid).link.status == "ATIVO"  # volta à folha dos ativos


def test_eventos_do_beneficio_exigem_senha(institution, master, setores, client_for, caso):
    from apps.concessao.models import Benefit
    from apps.institutions.models import PayrollGroup

    b = Benefit.objects.create(institution=institution, kind="APOSENTADORIA", beneficiary=caso["person"],
                               payroll_group=PayrollGroup.objects.get(institution=institution, segment="FINANCEIRO"),
                               start_date=date(2024, 1, 1), amount=Decimal("3000.00"))
    m = client_for(master)
    body = {"kind": "SUSPENSAO", "date": "2026-10-01", "justification": "Prova de vida não realizada no prazo."}
    assert m.post(f"/api/v1/beneficios/{b.pk}/event/", body, format="json").status_code == 409
    assert m.post(f"/api/v1/beneficios/{b.pk}/event/", {**body, "password": PASSWORD},
                  format="json").json()["status"] == "SUSPENSO"
